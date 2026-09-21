from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from gc_rpa_core.config import RpaConfig
from gc_rpa_core.db import DbEndpoint
from hd_seat_jit import __main__ as entry
from hd_seat_jit import sheet

LABELS = ("작업지시번호", "일 자", "공 장", "라 인", "순 번", "완성품사양", "수 량", "누 계")
BODY = ("20260921-001", "2026-09-21", 1, 3.0, 1.0, "ALC-A", 10.0, 120.0)
COMMON = ("2026-09-21 08:00", "조립1라인", "07:50", "42")


@pytest.fixture
def workbook() -> Callable[..., Path]:
    def build(
        path: Path,
        *,
        labels: tuple[str, ...] = LABELS,
        rows: tuple[tuple[Any, ...], ...] = (BODY,),
        common: tuple[str, str, str, str] = COMMON,
    ) -> Path:
        book = Workbook()
        page = book.active
        page["C4"], page["C7"], page["C9"], page["C10"] = common
        for column, label in enumerate(labels, start=1):
            page.cell(sheet.HEADER_ROW, column, label)
        for offset, values in enumerate(rows):
            for column, value in enumerate(values, start=1):
                page.cell(sheet.HEADER_ROW + 1 + offset, column, value)
        book.save(path)
        return path

    return build


def endpoint() -> DbEndpoint:
    return DbEndpoint(host="", port=None, database="", user="", password="")


@pytest.fixture
def rpa_settings(tmp_path: Path) -> RpaConfig:
    return RpaConfig(
        name="테스트 현대시트",
        url="https://test-site.invalid",
        user_id="test_user",
        password="test_pw",
        otp="",
        use_otp=False,
        move_path=str(tmp_path / "work"),
        exe_name="test_app.exe",
        source=endpoint(),
        target=endpoint(),
    )


@pytest.fixture(autouse=True)
def keep_real_browsers_alive(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(entry, "clean_up_browsers_on_exit", lambda: None)
