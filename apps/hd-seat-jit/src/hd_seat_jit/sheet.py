from __future__ import annotations

import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HEADER_ROW = 12
SPACES = re.compile(r"\s+")


class SheetError(RuntimeError):
    pass


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def number(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def line_text(value: Any) -> str:
    try:
        return str(int(float(value)))
    except (TypeError, ValueError):
        return text(value)


@dataclass(frozen=True)
class Column:
    name: str
    label: str
    read: Callable[[Any], Any]


COLUMNS = (
    Column("wo_no", "작업지시번호", text),
    Column("wo_dt", "일 자", text),
    Column("wo_fac", "공 장", number),
    Column("wo_line", "라 인", line_text),
    Column("wo_order", "순 번", number),
    Column("alc_cd", "완성품사양", text),
    Column("wo_qty", "수 량", number),
    Column("wo_sum", "누 계", number),
)

COMMON_CELLS = (
    ("wo_down_tm", "C4"),
    ("wo_line_nm", "C7"),
    ("crate_tm", "C9"),
    ("wo_cnt", "C10"),
)

CUSTOMER_COLUMN = "cust_cd"

FIELDS = (
    tuple(column.name for column in COLUMNS)
    + tuple(name for name, _ in COMMON_CELLS)
    + (CUSTOMER_COLUMN,)
)


@dataclass(frozen=True)
class Sheet:
    source: Path
    common: dict[str, str]
    rows: list[tuple[Any, ...]]


def squeezed(value: Any) -> str:
    return SPACES.sub("", text(value))


def header_columns(worksheet: Any) -> dict[str, int]:
    labels = {
        squeezed(worksheet.cell(HEADER_ROW, where).value): where
        for where in range(1, worksheet.max_column + 1)
    }
    placed = {}
    for column in COLUMNS:
        where = labels.get(squeezed(column.label))
        if where is None:
            raise SheetError(f"{HEADER_ROW}행에 '{column.label}' 열이 없습니다")
        placed[column.name] = where
    return placed


def body_values(worksheet: Any) -> Iterator[tuple[Any, ...]]:
    placed = header_columns(worksheet)
    for row in range(HEADER_ROW + 1, worksheet.max_row + 1):
        values = tuple(
            column.read(worksheet.cell(row, placed[column.name]).value) for column in COLUMNS
        )
        if values[0]:
            yield values


def read(path: Path, *, customer_code: str) -> Sheet:
    workbook = load_workbook(path, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        common = {name: text(worksheet[cell].value) for name, cell in COMMON_CELLS}
        trailing = (*(common[name] for name, _ in COMMON_CELLS), customer_code)
        rows = [(*values, *trailing) for values in body_values(worksheet)]
    finally:
        workbook.close()

    return Sheet(source=path, common=common, rows=rows)
